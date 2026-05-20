from flask import Flask, abort, render_template, request, redirect, url_for, session, send_file, jsonify, flash
from datetime import datetime
from markupsafe import Markup
from werkzeug.utils import secure_filename
import json, os
from openpyxl import Workbook
import io, uuid
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv
from collections import Counter
from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)
from datetime import timedelta
from functools import wraps
import cloudinary
import cloudinary.uploader

load_dotenv()

def format_date(date_str):
    return datetime.strptime(date_str, "%Y-%m-%d").strftime("%d-%m-%Y")

def admin_login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("admin_id"):
            flash(
                "Sesi Anda telah berakhir. Silakan login kembali.",
                "warning"
            )
            return redirect(url_for("admin_login"))
        return f(*args, **kwargs)
    return wrapper

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", os.urandom(32).hex())
app.config['SESSION_PERMANENT'] = True
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(minutes=15)
app.config["UPLOAD_FOLDER"] = "static/uploads"

app.jinja_env.globals.update(format_date=format_date)

app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=True
)

cloudinary.config(
    cloud_name=os.getenv("CLOUDINARY_CLOUD_NAME"),
    api_key=os.getenv("CLOUDINARY_API_KEY"),
    api_secret=os.getenv("CLOUDINARY_API_SECRET"),
    secure=True
)

from flask_sqlalchemy import SQLAlchemy

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DB_URL = os.getenv("DATABASE_URL")

app.config['SQLALCHEMY_DATABASE_URI'] = DB_URL


app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)

from flask_migrate import Migrate
migrate = Migrate(app, db)

ALLOWED_IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "webp"}

def allowed_image(filename):
    return (
        "." in filename and
        filename.rsplit(".", 1)[1].lower() in ALLOWED_IMAGE_EXTENSIONS
    )

def create_notification(message, notif_type="info", target_url=None):

    notif = Notification(
        message=message,
        type=notif_type,
        target_url=target_url,
        created_at=datetime.now().strftime("%Y-%m-%d %H:%M")
    )

    db.session.add(notif)
    db.session.commit()

class Survey(db.Model):
    id = db.Column(db.Integer, primary_key=True)

    title = db.Column(db.String(200), nullable=False)
    description = db.Column(db.Text)
    start_date = db.Column(db.String(20))
    end_date = db.Column(db.String(20))
    is_active = db.Column(db.Boolean, default=True)
    email_enabled = db.Column(db.Boolean, default=True)
    email_required = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.String(30))
    questions = db.Column(db.Text)
    unsur_responden = db.Column(db.Text)

    responses = db.relationship(
        'SurveyResponse',
        backref='survey',
        lazy=True,
        cascade="all, delete-orphan"
    )
    @property
    def total_responden(self):
        return len(self.responses)
    
    # =========================
    # 🔥 STATUS OTOMATIS SURVEY
    # =========================
    @property
    def status(self):
        if not self.start_date or not self.end_date:
            return "tidak_valid"

        now = datetime.now().date()

        start = datetime.strptime(
            self.start_date, "%Y-%m-%d"
        ).date()

        end = datetime.strptime(
            self.end_date, "%Y-%m-%d"
        ).date()

        if now < start:
            return "belum_mulai"
        elif start <= now <= end:
            return "aktif"
        else:
            return "tutup"

class SurveyResponse(db.Model):

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    survey_id = db.Column(
        db.Integer,
        db.ForeignKey('survey.id'),
        nullable=False
    )

    unsur_responden = db.Column(
        db.String(100)
    )

    email = db.Column(db.String(150))

    answers = db.Column(
        db.Text
    )

    created_at = db.Column(
        db.String(30)
    )

    @property
    def answers_json(self):
        return json.loads(self.answers or "{}")
    
class Notification(db.Model):

    id = db.Column(db.Integer, primary_key=True)
    message = db.Column(db.String(255))
    type = db.Column(db.String(50))

    target_url = db.Column(db.String(255)) 

    is_read = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.String(30))

class Admin(db.Model):

    id = db.Column(db.Integer, primary_key=True)

    username = db.Column(db.String(100))
    password = db.Column(db.String(255))

    photo = db.Column(
        db.String(255),
        default="/static/img/profile.png"
    )

    display_name = db.Column(
        db.String(100),
        default="Administrator"
    )

@app.context_processor
def inject_admin():

    admin = None

    if session.get("admin_id"):

        admin = Admin.query.get(
            session.get("admin_id")
        )

    return dict(admin=admin)

# -----------------------------
# LOGIN ADMIN
# -----------------------------
@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    error = None

    if request.method == "POST":

        username = request.form["username"]
        password = request.form["password"]

        admin = Admin.query.filter_by(
            username=username
        ).first()

        if admin and check_password_hash(admin.password, password):

            session.clear()
            session.permanent = True
            session["admin"] = True
            session["admin_id"] = admin.id

            return redirect(url_for("admin_dashboard"))

        else:
            error = "Username atau password salah"

    return render_template(
        "admin/login.html",
        error=error
    )

@app.route("/")
def home():
    return redirect(url_for("admin_login"))

@app.route("/dashboard")
@admin_login_required
def admin_dashboard():
    
    admin = Admin.query.get(
        session.get("admin_id")
    )

    surveys = Survey.query.all()
    total_survey = Survey.query.count()
    total_responden = SurveyResponse.query.count()
    survey_aktif = Survey.query.filter_by(
        is_active=True
    ).count()

    return render_template(
        "admin/dashboard.html",
        surveys=surveys,
        total_survey=total_survey,
        total_responden=total_responden,
        survey_aktif=survey_aktif,
        active_page="dashboard"
    )

@app.route('/dashboard/chart-data/<int:survey_id>/<int:q_index>')
def dashboard_chart_data(survey_id, q_index):

    survey = Survey.query.get_or_404(survey_id)
    questions = json.loads(survey.questions or "[]")
    target_question = None
    nomor = 1

    for section in questions:
        for q in section.get("questions", []):
            if nomor == q_index:
                target_question = q
                break
            nomor += 1

    if not target_question:
        return jsonify({
            "type": "empty"
        })

    # =========================
    # ISIAN
    # =========================
    if target_question.get("type") == "isian":

        return jsonify({
            "type": "text"
        })

    # =========================
    # PILIHAN
    # =========================

    responses = SurveyResponse.query.filter_by(
        survey_id=survey_id
    ).all()

    counter = Counter()
    answer_key = f"q{q_index}"
    label_map = {}

    for opt in target_question.get("options", []):
        label_map[opt["label"]] = opt["text"]

    for r in responses:
        answers = json.loads(r.answers or "{}")
        answer = answers.get(answer_key)

        if not answer:
            continue

        # MULTIPLE CHOICE
        if isinstance(answer, list):

            for item in answer:
                counter[item] += 1

        # SINGLE CHOICE
        else:
            counter[answer] += 1

    labels = []
    values = []

    for label, total in counter.items():

        option_text = label_map.get(label, label)

        labels.append(f"{label}. {option_text}")
        values.append(total)

    return jsonify({
        "type": "chart",
        "labels": labels,
        "values": values
    })

@app.route("/admin/notifications/unread-count")
@admin_login_required
def unread_count():

    if not session.get("admin"):
        return jsonify({"count": 0})

    count = Notification.query.filter_by(
        is_read=False
    ).count()

    return jsonify({
        "count": count
    })

@app.route("/admin/notifications/read", methods=["POST"])
@admin_login_required
def mark_all_notifications_read():
    if not session.get("admin"):
        return jsonify({"success": False})

    Notification.query.filter_by(is_read=False).update({
        "is_read": True
    })
    db.session.commit()
    return jsonify({"success": True})

@app.route("/survey")
@admin_login_required
def admin_survey():
    
    page = request.args.get("page", 1, type=int)
    per_page = request.args.get("per_page", 5, type=int)

    pagination = Survey.query.order_by(
        Survey.id.desc()
    ).paginate(
        page=page,
        per_page=per_page,
        error_out=False
    )

    return render_template(
        "admin/survey.html",
        surveys=pagination.items,
        pagination=pagination,
        per_page=per_page,
        active_page="survey"
    )

@app.route("/admin/notifikasi")
@admin_login_required
def get_notifications():
    notifs = Notification.query.order_by(
        Notification.id.desc()
    ).limit(10).all()

    return jsonify([
        {
            "id": n.id,
            "message": n.message,
            "type": n.type,
            "time": n.created_at,
            "is_read": n.is_read,
            "target_url": n.target_url
        }
        for n in notifs
    ])

@app.route("/admin/notifications/clear", methods=["POST"])
@admin_login_required
def clear_notifications():
    if not session.get("admin"):
        return jsonify({"success": False})

    Notification.query.delete()
    db.session.commit()

    return jsonify({"success": True})

@app.route("/admin/notifications/delete/<int:notif_id>", methods=["POST"])
@admin_login_required
def delete_notification(notif_id):
    if not session.get("admin"):
        return jsonify({"success": False})

    notif = Notification.query.get(notif_id)

    if notif:
        db.session.delete(notif)
        db.session.commit()

    return jsonify({"success": True})

@app.route("/upload-profile-photo", methods=["POST"])
@admin_login_required
def upload_profile_photo():

    if not session.get("admin"):
        return jsonify({"success": False}), 401

    file = request.files.get("photo")
    if not file or file.filename == "":
        return jsonify({"success": False, "error": "File kosong"}), 400

    if not allowed_image(file.filename):
        return jsonify({
            "success": False,
            "error": "Format file tidak didukung"
        }), 400

    # =========================
    # UPLOAD KE CLOUDINARY
    # =========================
    upload_result = cloudinary.uploader.upload(
        file,
        folder="profile_admin",
        public_id=f"profile_{session.get('admin_id')}",
        overwrite=True,
        resource_type="image"
    )

    photo_url = upload_result.get("secure_url")

    # update database
    admin = db.session.get(Admin, session.get("admin_id"))
    admin.photo = photo_url
    db.session.commit()

    return jsonify({
        "success": True,
        "photo_url": photo_url
    })

@app.route("/settings/update", methods=["POST"])
@admin_login_required
def update_admin_settings():

    if "admin_id" not in session:
        return redirect("/login")

    admin_id = session["admin_id"]

    username = request.form.get("username")
    password = request.form.get("password")
    confirm_password = request.form.get("confirm_password")

    # validasi password
    if password and password != confirm_password:

        flash("Konfirmasi password tidak cocok", "error")
        return redirect(request.referrer)

    # ambil admin
    admin = Admin.query.get(admin_id)

    if not admin:
        flash("Admin tidak ditemukan", "error")
        return redirect(request.referrer)

    # update username
    admin.username = username

    # update password jika diisi
    if password:
        admin.password = generate_password_hash(password)

    db.session.commit()

    flash("Pengaturan berhasil diperbarui", "success")

    return redirect(request.referrer)

@app.route("/admin/notification/<int:notif_id>")
@admin_login_required
def open_notification(notif_id):

    notif = Notification.query.get_or_404(notif_id)

    # tandai dibaca
    if not notif.is_read:
        notif.is_read = True
        db.session.commit()

    # redirect ke tujuan
    if notif.target_url:
        return redirect(notif.target_url)

    return redirect(url_for("admin_dashboard"))

@app.route("/admin/survey/tambah", methods=["GET", "POST"])
@admin_login_required
def admin_survey_tambah():

    if request.method == "POST":
        title = request.form.get("title")
        description = request.form.get("pembuka")
        start_date = request.form.get("start_date")
        end_date = request.form.get("end_date")
        unsur_responden = request.form.getlist("unsur_responden[]")
        email_enabled = request.form.get("email_enabled") == "1"
        email_required = request.form.get("email_required") == "1"

        # =========================
        # SESSION TITLE
        # =========================

        session_titles = request.form.getlist(
            "session_title[]"
        )

        session_descriptions = request.form.getlist(
            "session_description[]"
        )

        question_sessions = request.form.getlist(
            "question_session[]"
        )

        # container session
        sessions = []

        for i, s in enumerate(session_titles):

            sessions.append({
                "title": s,

                "description":
                    session_descriptions[i]
                    if i < len(session_descriptions)
                    else "",

                "questions": []
            })

        # =========================
        # QUESTIONS
        # =========================

        texts = request.form.getlist("pertanyaan[]")
        types = request.form.getlist("type[]")
        counts = request.form.getlist("count[]")
        required_flags = request.form.getlist("is_required[]")
        
        question_indexes = request.form.getlist(
            "question_index[]"
        )

        for i in range(len(texts)):

            real_index = question_indexes[i]

            mode = request.form.get(
                f"mode[{real_index}]",
                "single"
            )

            labels = request.form.getlist(
                f"option_label[{real_index}][]"
            )

            texts_opt = request.form.getlist(
                f"option_text[{real_index}][]"
            )

            options = []

            for j in range(len(texts_opt)):

                label = (
                    labels[j].strip()
                    if j < len(labels)
                    else ""
                )

                text = texts_opt[j].strip()

                # skip option kosong
                if not label and not text:
                    continue

                options.append({
                    "label": label,
                    "text": text
                })

            q = {
                "text": texts[i],

                "type":
                    types[i]
                    if i < len(types)
                    else "isian",

                "count":
                    int(counts[i])
                    if i < len(counts)
                    and counts[i].isdigit()
                    else 1,

                "mode": mode,

                "options": options,
                "is_required": True if i < len(required_flags) else False
            }

            # ambil session index
            session_index = int(
                question_sessions[i]
            )

            # masukkan ke session terkait
            sessions[session_index]["questions"].append(q)

        # 🔥 BARU BUAT OBJECT SURVEY
        new_survey = Survey(
            title=title,
            description=description,
            start_date=start_date,
            end_date=end_date,

            questions=json.dumps(sessions),
            unsur_responden=json.dumps(unsur_responden),

            email_enabled=email_enabled,
            email_required=email_required,

            is_active=True,
            created_at=datetime.now().strftime("%Y-%m-%d %H:%M")
        )

        db.session.add(new_survey)
        db.session.commit()

        return redirect(url_for("admin_survey"))

    return render_template("admin/survey_form.html", active_page="survey")

@app.route("/admin/survey/<int:survey_id>/toggle", methods=["POST"])
@admin_login_required
def toggle_survey_status(survey_id):

    survey = db.session.get(Survey, survey_id)

    if survey:
        survey.is_active = not survey.is_active
        db.session.commit()

    return redirect(url_for("admin_survey"))

@app.route("/admin/survey/<int:survey_id>/edit", methods=["GET", "POST"])
@admin_login_required
def edit_survey(survey_id):

    survey = Survey.query.get_or_404(survey_id)
    
    if request.method == "POST":

        # =========================
        # BASIC INFO
        # =========================

        survey.title = request.form.get("title")
        survey.description = request.form.get("pembuka")
        survey.start_date = request.form.get("start_date")
        survey.end_date = request.form.get("end_date")
        unsur_responden = request.form.getlist("unsur_responden[]")
        survey.unsur_responden = json.dumps(unsur_responden)
        required_flags = request.form.getlist("is_required[]")
        survey.email_enabled = request.form.get("email_enabled") == "1"
        survey.email_required = request.form.get("email_required") == "1"

        # =========================
        # SESSION TITLE
        # =========================

        session_titles = request.form.getlist(
            "session_title[]"
        )

        session_descriptions = request.form.getlist(
            "session_description[]"
        )

        question_sessions = request.form.getlist(
            "question_session[]"
        )

        sessions = []

        for i, s in enumerate(session_titles):

            sessions.append({
                "title": s,

                "description":
                    session_descriptions[i]
                    if i < len(session_descriptions)
                    else "",

                "questions": []
            })

        # =========================
        # QUESTIONS
        # =========================

        texts = request.form.getlist("pertanyaan[]")

        types = request.form.getlist("type[]")

        counts = request.form.getlist("count[]")

        question_indexes = request.form.getlist(
            "question_index[]"
        )

        for i in range(len(texts)):

            real_index = question_indexes[i]

            mode = request.form.get(
                f"mode[{real_index}]",
                "single"
            )

            labels = request.form.getlist(
                f"option_label[{real_index}][]"
            )

            texts_opt = request.form.getlist(
                f"option_text[{real_index}][]"
            )

            options = []

            for j in range(len(texts_opt)):

                label = (
                    labels[j].strip()
                    if j < len(labels)
                    else ""
                )

                text = texts_opt[j].strip()

                # skip option kosong
                if not label and not text:
                    continue

                options.append({
                    "label": label,
                    "text": text
                })

            q = {
                "text": texts[i],

                "type":
                    types[i]
                    if i < len(types)
                    else "isian",

                "count":
                    int(counts[i])
                    if i < len(counts)
                    and counts[i].isdigit()
                    else 1,

                "mode": mode,

                "options": options,
                "is_required": True if i < len(required_flags) else False
            }

            session_index = int(
                question_sessions[i]
            )

            sessions[session_index]["questions"].append(q)

        # =========================
        # SAVE QUESTIONS
        # =========================

        survey.questions = json.dumps(sessions)

        db.session.commit()

        return redirect(url_for("admin_survey"))

    return render_template(
        "admin/survey_form.html",
        survey=survey,
        active_page="survey",
        is_edit=True
    )


@app.route("/admin/survey/<int:survey_id>/detail")
@admin_login_required
def detail_survey(survey_id):
    
    response_id = request.args.get("response_id", type=int)

    survey = Survey.query.get_or_404(survey_id)

    responses = SurveyResponse.query.filter_by(
        survey_id=survey_id
    ).order_by(
        SurveyResponse.id.asc()
    ).all()

    parsed_questions = json.loads(
        survey.questions or "[]"
    )

    selected_response = None
    if response_id:
        selected_response = SurveyResponse.query.get(response_id)

    return render_template(
        "admin/detail_survey.html",
        survey=survey,
        responses=responses,
        parsed_questions=parsed_questions,
        selected_response=selected_response,
        active_page="survey"
    )


@app.post("/admin/respon/<int:response_id>/hapus")
@admin_login_required
def hapus_respon(response_id):

    response = SurveyResponse.query.get_or_404(response_id)

    db.session.delete(response)
    db.session.commit()

    return redirect(request.referrer or url_for("admin_survey"))

@app.route("/admin/survey/<int:survey_id>/hapus", methods=["POST"])
@admin_login_required
def hapus_survey(survey_id):

    survey = db.session.get(Survey, survey_id)

    if survey:
        db.session.delete(survey)
        db.session.commit()

    return redirect(url_for("admin_survey"))

@app.route("/survey/<int:survey_id>", methods=["GET", "POST"])
def survey_responden(survey_id):

    errors = []

    survey = Survey.query.get_or_404(survey_id)
    skip_intro = session.get(f"survey_{survey_id}_skip_intro", False)

    if f"survey_{survey_id}_skip_intro" not in session:
        unsur_list = json.loads(survey.unsur_responden or '["Semua"]')

        session[f"survey_{survey_id}_skip_intro"] = (
            (not survey.email_enabled)
            and (unsur_list == ["Semua"] or len(unsur_list) == 0)
        )

    # =========================
    # CEK STATUS SURVEY
    # =========================
    if survey.status != "aktif" or not survey.is_active:
        return render_template(
            "survey/survey_closed.html",
            survey=survey
        )

    # =========================
    # AMBIL TOKEN (GET / POST)
    # =========================
    token = (
        request.args.get("token")
        if request.method == "GET"
        else request.form.get("token")
    )

    session_token = session.get(
        f"survey_{survey_id}_token"
    )

    if not token or token != session_token:
        abort(404)
    
    survey = Survey.query.get_or_404(survey_id)

    # =========================
    # POST (SUBMIT JAWABAN)
    # =========================
    if request.method == "POST":

        token = request.form.get("token")
        session_token = session.get(f"survey_{survey_id}_token")
        email = None

        if survey.email_enabled:
            email = request.form.get("email", "").strip()

            if survey.email_required and not email:
                errors.append({
                    "number": "Email",
                    "text": "Email wajib diisi"
                })

            elif email:
                import re
                if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]{2,}$", email):
                    errors.append({
                        "number": "Email",
                        "text": "Format email tidak valid"
                    })

        if not token or token != session_token:
            abort(404)

        raw_sessions = json.loads(survey.questions or "[]")

        # format session
        if raw_sessions and "questions" not in raw_sessions[0]:
            sessions = [{
                "title": "Survey",
                "questions": raw_sessions
            }]
        else:
            sessions = raw_sessions

        # GLOBAL NUMBERING
        counter = 1
        for s in sessions:
            for q in s["questions"]:
                q["global_index"] = counter
                counter += 1

        unsur_responden = request.form.get("unsur_responden")
        errors = []
        answers = {}

        counter = 1
        error_session_index = 0

        for s_idx, session_item in enumerate(sessions):
            for q in session_item["questions"]:

                if q.get("is_required"):
                    key_prefix = f"q{counter}"

                    found = []

                    for key in request.form.keys():
                        if key.startswith(key_prefix):
                            found.extend(request.form.getlist(key))

                    # 🔥 jika ADA SATU SAJA kolom kosong → error
                    if not found or any(v.strip() == "" for v in found):
                        errors.append({
                            "number": counter,
                            "text": q["text"]
                        })
                        error_session_index = s_idx

                counter += 1

        # JIKA ADA ERROR
        if errors:
            pass

            unsur_list = json.loads(
                survey.unsur_responden or '["Semua"]'
            )

            return render_template(
                "survey/isi_survey.html",
                survey=survey,
                sessions=sessions,
                show_email=survey.email_enabled,
                email_required=survey.email_required,
                errors=errors,
                old=request.form,
                selected_unsur=unsur_responden,
                error_session_index=error_session_index,
                token=token,
                skip_intro=True,
                has_error=True,
            )

        # SIMPAN JAWABAN
        for key in request.form.keys():

            if key.startswith("q"):

                values = request.form.getlist(key)

                if len(values) == 1:
                    answers[key] = values[0]
                else:
                    answers[key] = values

        response = SurveyResponse(
            survey_id=survey_id,
            unsur_responden=unsur_responden,
            email=email if survey.email_enabled else None,
            answers=json.dumps(answers),
            created_at=datetime.now().strftime("%Y-%m-%d %H:%M")
        )

        db.session.add(response)
        db.session.commit()

        create_notification(
            f"Responden baru mengisi survey '{survey.title}'",
            "response",
            target_url=url_for("detail_survey", survey_id=survey.id, response_id=response.id)
        )
        session.pop(f"survey_{survey_id}_token", None)

        return render_template("survey/thank_you.html")

    raw_questions = json.loads(survey.questions or "[]")

    # =========================
    # FORMAT LAMA
    # =========================
    if raw_questions and "questions" not in raw_questions[0]:

        sessions = [
            {
                "title": "Survey",
                "questions": raw_questions
            }
        ]

    # =========================
    # FORMAT BARU
    # =========================
    else:
        sessions = raw_questions

    # =========================
    # GLOBAL NUMBERING
    # =========================
    counter = 1

    for session_item in sessions:

        for q in session_item["questions"]:

            q["global_index"] = counter
            counter += 1

    # 🔥 PARSE UNSUR
    unsur_responden = json.loads(
        survey.unsur_responden or '["Semua"]'
    )

    return render_template(
        "survey/isi_survey.html",
        survey=survey,
        sessions=sessions,
        unsur_responden=unsur_responden,
        token=request.args.get("token"),
        skip_intro=skip_intro
    )

@app.route("/survey/<int:survey_id>/pembuka")
def survey_intro(survey_id):

    survey = Survey.query.get_or_404(survey_id)

    # =========================
    # CEK STATUS SURVEY
    # =========================
    if survey.status != "aktif" or not survey.is_active:
        return render_template(
            "survey/survey_closed.html",
            survey=survey
        )

    return render_template(
        "survey/pembuka.html",
        survey=survey
    )

@app.route("/survey/<int:survey_id>/start", methods=["POST"])
def survey_start(survey_id):

    survey = Survey.query.get_or_404(survey_id)
    if not survey.is_active:
        abort(404)

    unsur_list = json.loads(survey.unsur_responden or '["Semua"]')

    skip_intro = (
        (not survey.email_enabled)
        and (unsur_list == ["Semua"] or len(unsur_list) == 0)
    )

    session[f"survey_{survey_id}_skip_intro"] = skip_intro

    # 🔐 token sekali pakai
    token = str(uuid.uuid4())

    session[f"survey_{survey_id}_token"] = token

    return redirect(
        url_for(
            "survey_responden",
            survey_id=survey_id,
            token=token
        )
    )

@app.template_filter('nl2br')
def nl2br(value):
    if not value:
        return ""
    return Markup(value.replace("\n", "<br>"))

@app.route("/admin/survey/<int:survey_id>/download")
@admin_login_required
def download_survey(survey_id):

    survey = Survey.query.get_or_404(
        survey_id
    )

    responses = SurveyResponse.query.filter_by(
        survey_id=survey_id
    ).all()

    wb = Workbook()

    ws = wb.active
    ws.title = "Respon Survey"

    # =====================================
    # HEADER
    # =====================================

    header_bottom = []

    sessions = json.loads(
        survey.questions or "[]"
    )

    question_keys = []

    # kolom awal
    base_headers = [
        ("No", True),
        ("Tanggal", True)
    ]

    # cek unsur responden
    unsur_list = json.loads(
        survey.unsur_responden or "[]"
    )

    show_unsur = (
        unsur_list
        and unsur_list != ["Semua"]
    )

    if show_unsur:
        base_headers.append(
            ("Unsur Responden", True)
        )

    # cek email
    if survey.email_enabled:
        base_headers.append(
            ("Email", True)
        )

    # tulis header awal
    for idx, (title, merge) in enumerate(base_headers, start=1):

        col_letter = get_column_letter(idx)

        if merge:
            ws.merge_cells(
                f"{col_letter}1:{col_letter}2"
            )

        ws.cell(
            row=1,
            column=idx
        ).value = title

    # posisi mulai pertanyaan
    current_col = len(base_headers) + 1

    counter = 1

    for session_item in sessions:

        questions = session_item.get(
            "questions",
            []
        )

        for q in questions:

            q_type = q.get("type")

            # =========================
            # ISIAN MULTI + RANKING
            # =========================

            if q_type in ["isian", "ranking"]:

                if q_type == "ranking":
                    count = len(q.get("options", []))
                else:
                    count = q.get("count", 1)

                start_col = current_col
                end_col = current_col + count - 1

                # merge judul pertanyaan
                ws.merge_cells(
                    start_row=1,
                    start_column=start_col,
                    end_row=1,
                    end_column=end_col
                )

                ws.cell(
                    row=1,
                    column=start_col
                ).value = f"{counter}. {q.get('text')}"

                for i in range(count):

                    if q_type == "ranking":
                        header_bottom.append(
                            f"Ranking {i+1}"
                        )
                    else:
                        header_bottom.append(
                            f"Jawaban {i+1}"
                        )

                    question_keys.append(
                        f"q{counter}"
                    )

                    current_col += 1

            # =========================
            # PILIHAN
            # =========================

            else:

                ws.merge_cells(
                    start_row=1,
                    start_column=current_col,
                    end_row=2,
                    end_column=current_col
                )

                ws.cell(
                    row=1,
                    column=current_col
                ).value = f"{counter}. {q.get('text')}"

                question_keys.append(
                    f"q{counter}"
                )

                header_bottom.append("")

                current_col += 1

            counter += 1


    # append header row 2
    for idx, value in enumerate(header_bottom):

        col = len(base_headers) + 1 + idx

        cell = ws.cell(row=2, column=col)

        # skip jika merged cell
        if cell.__class__.__name__ == "MergedCell":
            continue

        cell.value = value


    # =====================================
    # STYLE HEADER
    # =====================================

    for row in ws.iter_rows(
        min_row=1,
        max_row=2
    ):

        for cell in row:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

    # =====================================
    # DATA
    # =====================================

    for idx, r in enumerate(
        responses,
        start=1
    ):

        answer_data = json.loads(
            r.answers or "{}"
        )

        row = [
            idx,
            r.created_at
        ]

        if show_unsur:
            row.append(
                r.unsur_responden or ""
            )

        if survey.email_enabled:
            row.append(
                r.email or ""
            )

        counter = 1

        for session_item in sessions:

            for q in session_item.get(
                "questions",
                []
            ):

                key = f"q{counter}"

                value = answer_data.get(
                    key,
                    ""
                )

                # =====================
                # OPTION MAP
                # =====================

                option_map = {}

                for opt in q.get(
                    "options",
                    []
                ):

                    option_map[
                        opt.get("label")
                    ] = opt.get("text")

                # =====================
                # ISIAN
                # =====================

                if q.get("type") == "isian":

                    count = q.get("count", 1)

                    # MULTI INPUT
                    if count > 1:

                        for i in range(count):

                            multi_key = f"{key}_{i}"

                            row.append(
                                answer_data.get(multi_key, "")
                            )

                    # SINGLE INPUT
                    else:

                        row.append(
                            answer_data.get(key, "")
                        )

                # =====================
                # RANKING
                # =====================

                elif q.get("type") == "ranking":

                    ranking_result = []

                    for i in range(len(q.get("options", []))):

                        rank_key = f"{key}_rank_{i}"
                        text_key = f"{key}_text_{i}"

                        rank_value = answer_data.get(rank_key)
                        text_value = answer_data.get(text_key)

                        if rank_value and text_value:

                            ranking_result.append({
                                "rank": int(rank_value),
                                "text": text_value
                            })

                    ranking_result = sorted(
                        ranking_result,
                        key=lambda x: x["rank"]
                    )

                    total_ranking = len(q.get("options", []))

                    for i in range(total_ranking):

                        if i < len(ranking_result):

                            item = ranking_result[i]

                            row.append(
                                item["text"]
                            )

                        else:
                            row.append("")

                # =====================
                # MULTIPLE CHOICE
                # =====================

                elif isinstance(value, list):

                    hasil = []

                    for item in value:

                        teks = option_map.get(
                            item,
                            item
                        )

                        hasil.append(
                            f"{item}. {teks}"
                        )

                    row.append(
                        ", ".join(hasil)
                    )

                # =====================
                # SINGLE CHOICE
                # =====================

                else:

                    teks = option_map.get(
                        value,
                        value
                    )

                    if value:
                        row.append(
                            f"{value}. {teks}"
                        )
                    else:
                        row.append("")

                counter += 1

        ws.append(row)

    # =====================================
    # AUTO WIDTH
    # =====================================

    for col_idx in range(1, ws.max_column + 1):
        column_letter = get_column_letter(col_idx)
        length = 0

        for cell in ws[column_letter]:
            if cell.value:
                length = max(length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = length + 5

    # =====================================
    # SAVE
    # =====================================

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=
            f"survey_{survey.id}.xlsx",
        mimetype=
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route("/logout")
def logout():

    session.clear()

    return redirect(url_for("admin_login"))

if __name__ == "__main__":
    with app.app_context():

        admin = Admin.query.filter_by(
            username="admin"
        ).first()

        if admin:
            admin.password = generate_password_hash("admin123")
            db.session.commit()

            print("Password admin berhasil direset")

    app.run(host="0.0.0.0", port=int(os.getenv("PORT", 5000)))
